"""
One-time bootstrap script.
Reads the actual Physicals Master Excel and generates:
  1. apps/data-entry-real/config/form_config.json  — tabs/fields for the UI
  2. sql/nuton_load_daily_real_ddl.sql              — CREATE TABLE DDL
  3. sql/daily_report_view.sql                     — CREATE OR REPLACE VIEW

Run once:
  python scripts/generate_from_excel.py <path-to-excel>

After that, Excel is no longer needed. Edit form_config.json and sql/ files directly.
"""
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

EXCEL = sys.argv[1] if len(sys.argv) > 1 else "docs/Physicals_Master.xlsx"
REPO  = Path(__file__).parent.parent

CATALOG = "serverless_stable_82bi8w_catalog"
SCHEMA  = "data_entry"
TABLE   = "nuton_load_daily_real"
FULL    = f"{CATALOG}.{SCHEMA}.{TABLE}"
VIEW    = f"{CATALOG}.{SCHEMA}.daily_report"

# Columns auto-derived in Nuton Load (formula cells — not user input, not in UI)
AUTO_DERIVED_NL = {"DX", "EA", "HD"}

# Tab groupings: each entry is (tab_id, tab_label, [row1/row2 section keywords])
# Sections from the Excel are assigned to tabs by matching their row1 header
TAB_SECTIONS = [
    ("general",     "General",          ["General"]),
    ("ore",         "Ore Delivered",    ["Nuton", "Stacked", "Crushed", "Mining"]),
    ("raffinate",   "Raffinate",        ["Raffinate"]),
    ("pls",         "PLS",              ["PLS"]),
    ("plant",       "Plant",            ["Plant Raffinate", "Plant Feed", "Plant PLS", "Plant Organic", "Precip"]),
    ("electrolyte", "Electrolyte",      ["Lean Electrolyte", "Rich Electrolyte"]),
    ("ew_sx",       "EW / SX",          ["Electro-winning", "SX Extraction", "Rectifier"]),
    ("cathode",     "Cathode",          ["Copper Cathode"]),
    ("reagents",    "Reagents & Misc",  ["Organic", "Reagents", "Burro"]),
]

def to_snake(text: str) -> str:
    """Convert a display label to snake_case column name."""
    text = re.sub(r"[^\w\s]", "", str(text))
    text = re.sub(r"\s+", "_", text.strip().lower())
    text = re.sub(r"_+", "_", text)
    return text[:60]


def build_nl_col_map(ws):
    """
    Build mapping: col_letter -> {label, unit, section, subsection, snake}
    Headers span rows 1-4:
      row1 = top-level section
      row2 = subsection
      row3 = field name
      row4 = unit
    """
    # Track running section/subsection across merged cells
    current_r1 = ""
    current_r2 = ""

    col_map = {}
    for c in range(1, ws.max_column + 1):
        letter = get_column_letter(c)
        r1 = ws.cell(1, c).value
        r2 = ws.cell(2, c).value
        r3 = ws.cell(3, c).value
        r4 = ws.cell(4, c).value

        if r1 and str(r1).strip():
            current_r1 = str(r1).strip()
            current_r2 = ""
        if r2 and str(r2).strip():
            current_r2 = str(r2).strip()

        label_parts = [p for p in [r3, r4] if p and str(p).strip()]
        if not label_parts:
            continue

        field_name = str(r3).strip() if r3 else ""
        unit       = str(r4).strip() if r4 else ""
        full_label = field_name + (f" ({unit})" if unit else "")

        # build unique snake name: subsection_field or section_field
        prefix = current_r2 or current_r1 or ""
        snake  = to_snake(f"{prefix}_{field_name}" if prefix else field_name)

        col_map[letter] = {
            "col":        letter,
            "section":    current_r1,
            "subsection": current_r2,
            "label":      full_label,
            "field_name": field_name,
            "unit":       unit,
            "snake":      snake,
            "formula":    letter in AUTO_DERIVED_NL,
        }

    return col_map


def assign_tab(section: str, subsection: str) -> str:
    text = f"{section} {subsection}".lower()
    for tab_id, _, keywords in TAB_SECTIONS:
        for kw in keywords:
            if kw.lower() in text:
                return tab_id
    return "reagents"  # catch-all


def sql_type(unit: str) -> str:
    unit = (unit or "").lower()
    if any(x in unit for x in ["ton", "lb", "gal", "gpm", "g/l", "ppm", "amps", "ratio"]):
        return "DOUBLE"
    if "%" in unit:
        return "DOUBLE"
    if any(x in unit for x in ["date", "time"]):
        return "STRING"
    return "DOUBLE"  # default numeric


def make_snake_unique(snake: str, seen: dict) -> str:
    base = snake
    n = seen.get(base, 0)
    seen[base] = n + 1
    return base if n == 0 else f"{base}_{n}"


def parse_dr_formula(formula: str, nl_col_map: dict, snake_by_letter: dict) -> str | None:
    """
    Translate a Daily Report INDEX/MATCH formula to a SQL column reference.
    Returns SQL expression string or None if not translatable.
    """
    if not formula or not str(formula).startswith("="):
        return None

    formula_str = str(formula)

    # SUM(X7:Y7) — two DR columns added together
    sum_match = re.search(r"SUM\(([A-Z]+)7:([A-Z]+)7\)", formula_str)
    if sum_match:
        c1 = sum_match.group(1)
        c2 = sum_match.group(2)
        # These are DR column refs — we need their NL source
        # They'll be resolved after all DR columns are mapped
        return f"__SUM__{c1}__{c2}__"

    # INDEX('Nuton Load'!XX:XX, MATCH(...))
    nl_refs = re.findall(r"'Nuton Load'!\$?([A-Z]+):\$?[A-Z]+", formula_str)
    if not nl_refs:
        return None

    # First ref is the data column, second is the match key (date col — skip)
    data_col = nl_refs[0]
    snake = snake_by_letter.get(data_col)
    if not snake:
        return None

    # Multiply by 100 for % efficiency
    if "*100" in formula_str or "* 100" in formula_str:
        return f"`{snake}` * 100"

    return f"`{snake}`"


def main():
    print(f"Loading {EXCEL}...")
    wb = load_workbook(EXCEL, data_only=False)

    nl_ws = wb["Nuton Load"]
    dr_ws = wb["Daily Report"]

    # ── 1. Build Nuton Load column map ────────────────────────────────────────
    print("Parsing Nuton Load headers...")
    nl_col_map = build_nl_col_map(nl_ws)

    # Deduplicate snake names
    seen_snakes: dict = {}
    snake_by_letter: dict = {}  # letter -> final snake name

    for letter, info in nl_col_map.items():
        unique = make_snake_unique(info["snake"], seen_snakes)
        nl_col_map[letter]["snake"] = unique
        snake_by_letter[letter] = unique

    # ── 2. Build form_config.json ─────────────────────────────────────────────
    print("Building form_config.json...")

    tabs = {tab_id: {"id": tab_id, "label": label, "fields": []}
            for tab_id, label, _ in TAB_SECTIONS}

    # Always add stacking_date to General tab as date picker
    tabs["general"]["fields"].append({
        "snake":  "stacking_date",
        "label":  "Stacking Date",
        "unit":   "",
        "type":   "date",
        "col":    "B",
    })

    for letter, info in nl_col_map.items():
        if letter in {"A", "B"}:  # date cols — handled separately
            continue
        # skip cols with no real label
        if not info["field_name"] or info["snake"].startswith("_"):
            continue
        if info["formula"]:  # auto-derived — not in UI
            continue
        if not info["field_name"]:
            continue

        tab_id = assign_tab(info["section"], info["subsection"])
        if tab_id not in tabs:
            tab_id = "reagents"

        tabs[tab_id]["fields"].append({
            "snake":      info["snake"],
            "label":      info["label"],
            "unit":       info["unit"],
            "type":       "number",
            "col":        letter,
            "section":    info["section"],
            "subsection": info["subsection"],
        })

    form_config = {
        "table":    FULL,
        "join_key": "stacking_date",
        "tabs":     [tabs[t[0]] for t in TAB_SECTIONS if tabs[t[0]]["fields"]],
    }

    config_path = REPO / "apps/data-entry-real/config/form_config.json"
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(json.dumps(form_config, indent=2))
    print(f"  → {config_path}  ({sum(len(t['fields']) for t in form_config['tabs'])} fields across {len(form_config['tabs'])} tabs)")

    # ── 3. Generate DDL ───────────────────────────────────────────────────────
    print("Generating DDL...")
    ddl_lines = [
        f"CREATE SCHEMA IF NOT EXISTS {CATALOG}.{SCHEMA};\n",
        f"CREATE TABLE IF NOT EXISTS {FULL} (",
        "    stacking_date  STRING NOT NULL,",
    ]

    for letter, info in nl_col_map.items():
        if letter in {"A", "B"}:
            continue
        if not info["field_name"] or info["snake"].startswith("_"):
            continue
        col_type = sql_type(info["unit"])
        ddl_lines.append(f"    `{info['snake']}`  {col_type},")

    # audit cols
    ddl_lines += [
        "    _submitted_by  STRING,",
        "    _submitted_at  STRING,",
        "    _tab_saves     STRING",   # JSON map of which tabs are saved
        ")",
        "USING DELTA",
        "COMMENT 'Nuton Load daily operator entry — generated from Physicals Master Excel';",
    ]

    ddl_sql = "\n".join(ddl_lines)
    ddl_path = REPO / "sql/nuton_load_daily_real_ddl.sql"
    ddl_path.parent.mkdir(parents=True, exist_ok=True)
    ddl_path.write_text(ddl_sql)
    print(f"  → {ddl_path}")

    # ── 4. Parse Daily Report formulas → view SQL ─────────────────────────────
    print("Parsing Daily Report formulas...")

    # Build DR col letter -> (label, formula)
    dr_cols = []
    for c in range(1, dr_ws.max_column + 1):
        letter = get_column_letter(c)
        formula = dr_ws.cell(7, c).value
        if not formula:
            continue

        # Build DR column label from rows 3-6
        parts = []
        for r in range(3, 7):
            v = dr_ws.cell(r, c).value
            if v and str(v).strip():
                parts.append(str(v).strip())
        label = " > ".join(parts)
        snake = to_snake(label)

        sql_expr = parse_dr_formula(formula, nl_col_map, snake_by_letter)
        dr_cols.append({
            "dr_col":  letter,
            "label":   label,
            "snake":   snake,
            "formula": str(formula),
            "sql":     sql_expr,
        })

    # Build DR letter -> sql_expr map for resolving SUM refs
    dr_letter_to_sql = {d["dr_col"]: d["sql"] for d in dr_cols}

    # Resolve __SUM__ placeholders
    select_parts = ["    stacking_date"]
    seen_view_cols: dict = {}
    for d in dr_cols:
        sql = d["sql"]
        if sql and sql.startswith("__SUM__"):
            parts = sql.replace("__SUM__", "").split("__")
            parts = [p for p in parts if p]
            resolved = []
            for dr_ref in parts:
                ref_sql = dr_letter_to_sql.get(dr_ref)
                if ref_sql and not ref_sql.startswith("__SUM__"):
                    resolved.append(f"COALESCE({ref_sql}, 0)")
                else:
                    resolved.append("0")
            sql = " + ".join(resolved)

        if sql:
            col_snake = make_snake_unique(to_snake(d["label"]), seen_view_cols)
            select_parts.append(f"    {sql} AS `{col_snake}`")

    view_sql = (
        f"CREATE OR REPLACE VIEW {VIEW} AS\n"
        f"SELECT\n"
        + ",\n".join(select_parts) + "\n"
        f"FROM {FULL};\n"
    )

    view_path = REPO / "sql/daily_report_view.sql"
    view_path.write_text(view_sql)
    print(f"  → {view_path}  ({len(select_parts)-1} columns)")

    print("\nDone. Files generated:")
    print(f"  {config_path}")
    print(f"  {ddl_path}")
    print(f"  {view_path}")
    print("\nNext: review the files, commit them, then scaffold apps/data-entry-real/")


if __name__ == "__main__":
    main()
